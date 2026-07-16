from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import get_settings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def parse_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    raise ValueError(f"unsupported file type: {suffix}")


def chunk_text(text: str) -> list[dict]:
    settings = get_settings()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", "。", ".", "；", ";", "，", ",", " ", ""],
    )
    pieces = [piece.strip() for piece in splitter.split_text(text) if piece.strip()]
    return [
        {
            "content": piece,
            "metadata": {"chunk_index": index, "chunk_total": len(pieces)},
        }
        for index, piece in enumerate(pieces)
    ]


def _parse_pdf(path: Path) -> str:
    parts: list[str] = []
    with pymupdf.open(path) as document:
        for page in document:
            text = page.get_text().strip()
            if text:
                parts.append(text)
            for table in page.find_tables():
                rows = table.extract()
                if rows:
                    parts.append(_rows_to_markdown(rows))
    return "\n\n".join(parts)


def _parse_docx(path: Path) -> str:
    document = DocxDocument(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            parts.append(_rows_to_markdown(rows))
    return "\n\n".join(parts)


def _rows_to_markdown(rows: list[list]) -> str:
    normalized = [[str(cell or "") for cell in row] for row in rows]
    if not normalized:
        return ""
    output = ["| " + " | ".join(normalized[0]) + " |"]
    output.append("| " + " | ".join(["---"] * len(normalized[0])) + " |")
    output.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(output)
